import io
import zipfile

import pytest
from docx import Document
from reportlab.pdfgen.canvas import Canvas

from mindchime.parsers import parse_document


def test_text_markdown_encoding_and_empty():
    parsed = parse_document("# 方法\n先验证\n# 边界\n不盲目行动".encode(), "资料.md")
    assert len(parsed["chapters"]) == 2
    assert parsed["chapters"][1]["title"] == "边界"
    gbk = parse_document("先检查成本".encode("gb18030"), "资料.txt")
    assert "成本" in gbk["chapters"][0]["text"]
    with pytest.raises(ValueError, match="为空"):
        parse_document(b"", "empty.txt")


def test_markdown_h4_h5_h6_are_real_sections():
    parsed = parse_document("# 书名\n作者说明\n#### 第一章：先缩小未知\n原文A\n##### 第二章：方法\n原文B\n###### 第三节：边界\n原文C".encode(), "资料.md")
    assert [c["title"] for c in parsed["chapters"]] == ["书名", "第一章：先缩小未知", "第二章：方法", "第三节：边界"]
    assert parsed["chapters"][1]["text"].endswith("原文A\n")


def test_pdf_text_pages_and_scanned_page_flag():
    raw = io.BytesIO()
    canvas = Canvas(raw)
    canvas.drawString(30, 700, "A real page with text")
    canvas.showPage()
    canvas.showPage()
    canvas.save()
    parsed = parse_document(raw.getvalue(), "pages.pdf")
    assert parsed["chapters"][0]["locator"] == {"type": "pdf_page", "page": 1}
    assert "real page" in parsed["chapters"][0]["text"]
    assert parsed["chapters"][1]["needs_ocr"] is True
    assert parsed["warnings"]


def test_docx_headings_and_table():
    doc = Document()
    doc.add_heading("验证", 1)
    doc.add_paragraph("检查方法的适用条件")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "输入"
    table.rows[0].cells[1].text = "输出"
    raw = io.BytesIO()
    doc.save(raw)
    parsed = parse_document(raw.getvalue(), "方法.docx")
    assert any("适用条件" in c["text"] for c in parsed["chapters"])
    assert any(c["locator"]["type"] == "docx_table" and "输入" in c["text"] for c in parsed["chapters"])


def test_epub_uses_spine_order_and_document_locator():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("META-INF/container.xml", '<container><rootfiles><rootfile full-path="EPUB/book.opf"/></rootfiles></container>')
        archive.writestr("EPUB/book.opf", '<package><manifest><item id="a" href="a.xhtml" media-type="application/xhtml+xml"/><item id="b" href="b.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="b"/><itemref idref="a"/></spine></package>')
        archive.writestr("EPUB/a.xhtml", '<html><h1>甲</h1><p>第一份内容</p></html>')
        archive.writestr("EPUB/b.xhtml", '<html><h1>乙</h1><p>第二份内容</p><script>danger()</script></html>')
    parsed = parse_document(output.getvalue(), "demo.epub")
    assert [c["title"] for c in parsed["chapters"]] == ["乙", "甲"]
    assert parsed["chapters"][0]["locator"]["path"] == "EPUB/b.xhtml"
    assert "danger" not in parsed["chapters"][0]["text"]


@pytest.mark.parametrize("name,raw", [("broken.pdf", b"bad"), ("broken.epub", b"bad"), ("movie.mp4", b"bad")])
def test_corrupt_and_deferred_formats_are_explicit(name, raw):
    with pytest.raises(ValueError):
        parse_document(raw, name)
