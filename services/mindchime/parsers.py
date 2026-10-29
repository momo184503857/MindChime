"""Offline document parsing. Every section retains an honest source locator."""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

from bs4 import BeautifulSoup
from defusedxml import ElementTree as ET
from docx import Document
from pypdf import PdfReader


SUPPORTED = {".pdf", ".epub", ".txt", ".md", ".markdown", ".docx"}


def section(title, text, locator, *, needs_ocr=False):
    return {"title": title, "text": text, "locator": locator, "selected": bool(text.strip()),
            "needs_ocr": needs_ocr}


def decode_text(raw: bytes) -> str:
    encodings = ["utf-16", "utf-8-sig", "gb18030"] if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else ["utf-8-sig", "gb18030"]
    for encoding in encodings:
        try:
            text = raw.decode(encoding)
            if "\x00" not in text:
                return text.replace("\r\n", "\n").replace("\r", "\n")
        except (UnicodeDecodeError, UnicodeError):
            pass
    raise ValueError("无法识别文本编码，请保存为 UTF-8 后重试。")


def text_sections(text: str, markdown=False) -> list[dict]:
    if not text.strip():
        raise ValueError("文件没有可读取文字。")
    headings = list(re.finditer(r"(?m)^#{1,6}\s+(.+)$" if markdown else r"(?m)^(第[一二三四五六七八九十百千万零\d]+[章节部篇][^\n]{0,90})$", text))
    parts = []
    if headings:
        if text[:headings[0].start()].strip():
            parts.append(section("前言与导入", text[:headings[0].start()], {"type": "text_offsets", "start": 0, "end": headings[0].start()}))
        for i, match in enumerate(headings):
            end = headings[i+1].start() if i+1 < len(headings) else len(text)
            parts.append(section(match.group(1).strip(), text[match.start():end], {"type": "text_offsets", "start": match.start(), "end": end}))
    else:
        start = 0
        while start < len(text):
            end = min(start + 18000, len(text))
            if end < len(text):
                boundary = text.rfind("\n", start + 9000, end)
                if boundary > start:
                    end = boundary + 1
            parts.append(section(f"分段 {len(parts)+1}", text[start:end], {"type": "text_offsets", "start": start, "end": end}))
            start = end
    return parts


def parse_document(raw: bytes, filename: str) -> dict:
    if not raw:
        raise ValueError("文件为空，请选择有效资料。")
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError("支持 PDF、EPUB、TXT、Markdown 和 DOCX；网页与音视频导入暂未开放。")
    warnings = []
    try:
        if suffix in {".txt", ".md", ".markdown"}:
            chapters = text_sections(decode_text(raw), suffix != ".txt")
        elif suffix == ".pdf":
            reader = PdfReader(io.BytesIO(raw), strict=False)
            if reader.is_encrypted:
                raise ValueError("文件受密码保护，请使用可读取版本；本产品不解除文件保护。")
            chapters = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                chapters.append(section(f"文档第 {i+1} 页", text, {"type": "pdf_page", "page": i+1}, needs_ocr=not text.strip()))
            if any(x["needs_ocr"] for x in chapters):
                warnings.append("部分页面未提取到文字。可使用 OCR 识别后核对；显示的是文档页序号，不是印刷页码。")
        elif suffix == ".docx":
            doc = Document(io.BytesIO(raw))
            chapters = []
            title, paragraphs, start = "分段 1", [], 1
            for number, para in enumerate(doc.paragraphs, 1):
                if para.style and para.style.name.startswith("Heading") and paragraphs:
                    chapters.append(section(title, "\n".join(paragraphs), {"type": "docx_paragraph", "start": start, "end": number-1}))
                    title, paragraphs, start = para.text or f"分段 {len(chapters)+1}", [], number
                paragraphs.append(para.text)
            if paragraphs:
                chapters.append(section(title, "\n".join(paragraphs), {"type": "docx_paragraph", "start": start, "end": len(doc.paragraphs)}))
            for i, table in enumerate(doc.tables, 1):
                text = "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
                chapters.append(section(f"表格 {i}", text, {"type": "docx_table", "table": i}))
        else:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 250_000_000:
                    raise ValueError("EPUB 解压后内容过大，请拆分资料。")
                if "META-INF/encryption.xml" in archive.namelist():
                    encryption = archive.read("META-INF/encryption.xml")
                    if b"EncryptedData" in encryption and b"font" not in encryption.lower():
                        raise ValueError("EPUB 含受保护内容，请提供可读取版本。")
                root = ET.fromstring(archive.read("META-INF/container.xml"))
                rootfile = next((n.attrib.get("full-path") for n in root.iter() if n.tag.endswith("rootfile")), None)
                if not rootfile:
                    raise ValueError("EPUB 缺少内容目录。")
                package = ET.fromstring(archive.read(rootfile))
                manifest = {n.attrib["id"]: n.attrib for n in package.iter() if n.tag.endswith("item") and "id" in n.attrib}
                ordered = [manifest[n.attrib["idref"]] for n in package.iter() if n.tag.endswith("itemref") and n.attrib.get("idref") in manifest]
                chapters = []
                import posixpath
                from urllib.parse import unquote
                for item in ordered:
                    if "html" not in item.get("media-type", ""):
                        continue
                    location = posixpath.normpath(posixpath.join(posixpath.dirname(rootfile), unquote(item["href"].split("#")[0])))
                    if location.startswith("../") or location.startswith("/"):
                        raise ValueError("EPUB 内容路径无效。")
                    soup = BeautifulSoup(archive.read(location), "html.parser")
                    for tag in soup(["script", "style", "nav"]):
                        tag.decompose()
                    title_tag = soup.find(["h1", "h2", "h3", "title"])
                    title = title_tag.get_text(" ", strip=True) if title_tag else f"分段 {len(chapters)+1}"
                    text = soup.get_text("\n", strip=True)
                    chapters.append(section(title, text, {"type": "epub_document", "path": location}))
        if not chapters:
            raise ValueError("文件未包含可读取章节。")
        count = sum(len(re.sub(r"\s", "", x["text"])) for x in chapters)
        if not count and suffix != ".pdf":
            raise ValueError("文件没有可读取文字。")
        return {"chapters": chapters, "warnings": warnings, "characters": count,
                "status": "needs_ocr" if not count else "parsed", "format": suffix[1:]}
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("文件已损坏或无法解析，请检查文件后重试。") from exc


def pdf_page_image(path: str, page: int) -> str:
    import base64
    import pypdfium2 as pdfium
    document = pdfium.PdfDocument(path)
    try:
        image = document[page-1].render(scale=1.5).to_pil()
        output = io.BytesIO()
        image.save(output, format="PNG")
        return "data:image/png;base64," + base64.b64encode(output.getvalue()).decode()
    finally:
        document.close()
