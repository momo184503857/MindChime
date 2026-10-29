"""Create original, labelled fixtures for parser and OCR acceptance."""
import zipfile
from pathlib import Path

import pypdfium2
from docx import Document
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / ".runtime/fixtures"
OUTPUT.mkdir(parents=True, exist_ok=True)
TEXT = ["MindChime 原创测试资料（非研究结论）", "第一章：小步验证", "先选择一个可逆的小动作，限定十五分钟，记录真实结果。",
        "这只是演示方法，不保证项目成功。", "第二章：检查边界", "涉及他人承诺或不可逆损失时，先澄清目标和代价。"]
pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
pdf_path = OUTPUT / "original-text.pdf"
pdf = canvas.Canvas(str(pdf_path))
for page in (TEXT[:4], [TEXT[0], *TEXT[4:]]):
    pdf.setFont("STSong-Light", 16)
    y = 760
    for line in page:
        pdf.drawString(50, y, line)
        y -= 45
    pdf.showPage()
pdf.save()
scanned = canvas.Canvas(str(OUTPUT / "original-scan.pdf"))
document = pypdfium2.PdfDocument(str(pdf_path))
for page in document:
    rendered = page.render(scale=2).to_pil()
    scanned.drawImage(ImageReader(rendered), 0, 0, width=595.28, height=841.89)
    scanned.showPage()
    page.close()
scanned.save()
document.close()
docx = Document()
docx.add_heading(TEXT[0], 0)
for line in TEXT[1:]:
    if line.startswith(("第一章", "第二章")):
        docx.add_heading(line, 1)
    else:
        docx.add_paragraph(line)
docx.save(OUTPUT / "original.docx")
with zipfile.ZipFile(OUTPUT / "original.epub", "w") as archive:
    archive.writestr("mimetype", "application/epub+zip")
    archive.writestr("META-INF/container.xml", '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
    archive.writestr("OEBPS/content.opf", '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">mindchime-test</dc:identifier><dc:title>原创测试</dc:title><dc:language>zh</dc:language></metadata><manifest><item id="one" href="one.xhtml" media-type="application/xhtml+xml"/><item id="two" href="two.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="one"/><itemref idref="two"/></spine></package>')
    for name, lines in [("one", TEXT[:4]), ("two", [TEXT[0], *TEXT[4:]])]:
        archive.writestr(f"OEBPS/{name}.xhtml", '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>原创测试</title></head><body>' + "".join(f"<p>{line}</p>" for line in lines) + "</body></html>")
(OUTPUT / "empty.txt").write_text("")
(OUTPUT / "corrupt.pdf").write_bytes(b"this is not a PDF")
print("Created original PDF, scanned PDF, DOCX, EPUB and invalid fixtures in .runtime/fixtures")
